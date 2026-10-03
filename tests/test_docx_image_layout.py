from __future__ import annotations

import base64
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import fitz
from docx import Document
from docx.shared import Pt

from app.docx_exporter import (
    _add_layout_text,
    _add_layout_picture,
    _export_facsimile_docx,
    _find_pandoc,
    _picture_rect_points,
    _referenced_picture_indices,
)


PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+/lZkAAAAASUVORK5CYII="
)


class DocxImageLayoutTests(unittest.TestCase):
    def test_unreadable_formula_uses_source_pdf_crop(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            source = Path(temporary) / "source.pdf"
            pdf = fitz.open()
            pdf.new_page(width=300, height=400)
            pdf.save(source)
            pdf.close()
            document = Document()
            with patch("app.docx_exporter._latex_to_omml", return_value=None):
                _add_layout_text(
                    document,
                    {"text": r"\frac{", "label": "formula",
                     "bbox": {"l": 40, "t": 350, "r": 140, "b": 320,
                              "coord_origin": "BOTTOMLEFT"}},
                    300, 400, source_pdf_path=source, page_number=1,
                )
            paragraph = document.paragraphs[0]
            self.assertEqual(paragraph.text, "")
            image = paragraph._p.xpath(".//wp:inline/wp:docPr")
            self.assertEqual(len(image), 1)
            self.assertEqual(image[0].get("descr"), "source-page=1;OCR-formula-fallback")

    @unittest.skipUnless(_find_pandoc(), "Pandoc is required for editable Word equations")
    def test_layout_option_keeps_marker_equation_and_page_position(self) -> None:
        document = Document()
        _add_layout_text(
            document,
            {"text": r"$ |\Delta E_{p1}| < |\Delta E_{p2}| $", "label": "list_item",
             "marker": "A.", "bbox": {"l": 57, "t": 314, "r": 147, "b": 302,
                                      "coord_origin": "BOTTOMLEFT"}},
            595, 841,
        )
        paragraph = document.paragraphs[0]
        self.assertEqual(paragraph.text, "A. ")
        self.assertEqual(len(paragraph._p.xpath(".//m:oMath")), 1)
        self.assertEqual(len(paragraph._p.xpath(".//w:framePr")), 1)
        frame = paragraph._p.xpath(".//w:framePr")[0]
        self.assertEqual(frame.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}x"), "1140")
        self.assertEqual(frame.get("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}y"), "10540")
        self.assertNotIn("$$", paragraph.text)

    def test_bottom_left_pdf_box_converts_to_page_coordinates(self) -> None:
        bbox = {"l": 100, "t": 740, "r": 160, "b": 690, "coord_origin": "BOTTOMLEFT"}
        self.assertEqual(_picture_rect_points(bbox, 595, 841), (100, 101, 160, 151))

    def test_removed_picture_is_not_exported(self) -> None:
        data = {"body": {"children": [{"$ref": "#/pictures/0"}]},
                "pictures": [{"children": []}, {"children": []}]}
        self.assertEqual(_referenced_picture_indices(data), {0})

    def test_picture_is_anchored_to_page_at_pdf_position(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            image = Path(temporary) / "tiny.png"
            image.write_bytes(PNG)
            document = Document()
            paragraph = document.add_paragraph()
            element = {"image_path": image, "page": 3, "index": 7,
                       "bbox": {"l": 100, "t": 740, "r": 160, "b": 690,
                                "coord_origin": "BOTTOMLEFT"}}
            _add_layout_picture(paragraph, element, 595, 841)
            anchors = document.element.xpath(".//wp:anchor")
            self.assertEqual(len(anchors), 1)
            anchor = anchors[0]
            self.assertEqual(anchor.xpath("./wp:positionH/wp:posOffset")[0].text, str(Pt(100)))
            self.assertEqual(anchor.xpath("./wp:positionV/wp:posOffset")[0].text, str(Pt(101)))
            self.assertEqual(anchor.xpath("./wp:positionH")[0].get("relativeFrom"), "page")
            self.assertEqual(anchor.xpath("./wp:docPr")[0].get("descr"), "source-page=3;picture-index=7")

    def test_facsimile_keeps_a_full_page_image_on_each_source_page(self) -> None:
        with tempfile.TemporaryDirectory() as temporary:
            directory = Path(temporary)
            source = directory / "source.pdf"
            pdf = fitz.open()
            pdf.new_page(width=300, height=400)
            pdf.new_page(width=300, height=400)
            pdf.save(source)
            pdf.close()
            output = _export_facsimile_docx(directory, source)
            document = Document(output)
            anchors = document.element.xpath(".//wp:anchor")
            self.assertEqual(len(anchors), 2)
            self.assertTrue(all(anchor.get("behindDoc") == "1" for anchor in anchors))
            self.assertEqual(len(document.element.xpath('.//w:br[@w:type="page"]')), 1)
            hybrid = _export_facsimile_docx(
                directory, source, mode="hybrid", page_texts={1: "可编辑 OCR 文本"}
            )
            hybrid_document = Document(hybrid)
            self.assertEqual(len(hybrid_document.element.xpath(".//wp:anchor")), 2)
            self.assertTrue(any("可编辑 OCR 文本" in p.text for p in hybrid_document.paragraphs))


if __name__ == "__main__":
    unittest.main()
