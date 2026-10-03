from __future__ import annotations

import html
import json
import math
import re
import shutil
import subprocess
import tempfile
from copy import deepcopy
from pathlib import Path
from urllib.parse import unquote, urlparse
from zipfile import ZipFile

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from lxml import etree

from app.dataset_validator import validate_dataset


DOCX_EXPORT_MODES = {"editable", "facsimile", "hybrid", "layout_editable"}
EMUS_PER_POINT = 12700
PDF_RENDER_DPI = 180
OMML_NAMESPACE = "http://schemas.openxmlformats.org/officeDocument/2006/math"
IMAGE_PATTERN = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
HEADING_PATTERN = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
BULLET_PATTERN = re.compile(r"^\s*[-*+]\s+(.+?)\s*$")
NUMBERED_PATTERN = re.compile(r"^\s*\d+[.)、]\s+(.+?)\s*$")
FORMULA_DELIMITER = "$$"


def export_dataset_to_docx(
    dataset_path: Path,
    mode: str = "editable",
    pdf_path: Path | None = None,
) -> Path:
    mode = mode or "editable"
    if mode not in DOCX_EXPORT_MODES:
        raise ValueError("Word 导出模式无效，请选择 editable、facsimile、hybrid 或 layout_editable。")

    raw_dataset_path = dataset_path.expanduser().resolve()
    dataset_dir = _resolve_dataset_dir(raw_dataset_path)
    export_base_dir = dataset_dir if dataset_dir is not None else raw_dataset_path.parent

    if mode == "editable":
        if dataset_dir is None:
            raise ValueError("可编辑文本模式需要提供数据集目录。")
        return _export_editable_docx(dataset_dir, dataset_dir, mode)

    if mode == "layout_editable":
        if dataset_dir is None:
            raise ValueError("可编辑版式模式需要提供数据集目录。")
        source_pdf_path = _resolve_pdf_path(raw_dataset_path, dataset_dir, pdf_path)
        return _export_layout_editable_docx(dataset_dir, dataset_dir, source_pdf_path)

    resolved_pdf_path = _resolve_pdf_path(raw_dataset_path, dataset_dir, pdf_path)
    if resolved_pdf_path is None:
        raise ValueError("未找到原始 PDF。请在请求中提供有效的 pdf_path。")

    page_texts = _extract_page_texts(dataset_dir) if mode == "hybrid" and dataset_dir is not None else {}
    fallback_markdown = _read_markdown_text(dataset_dir) if mode == "hybrid" and dataset_dir is not None else ""
    return _export_pdf_layout_docx(
        export_base_dir=export_base_dir,
        source_pdf_path=resolved_pdf_path,
        mode=mode,
        page_texts=page_texts,
        fallback_markdown=fallback_markdown,
    )


def _export_editable_docx(dataset_dir: Path, export_base_dir: Path, mode: str) -> Path:
    validation = validate_dataset(dataset_dir)
    if not validation.valid:
        detail = "；".join(validation.errors) or "数据集校验失败。"
        raise ValueError(detail)
    if validation.markdown_file is None:
        raise ValueError("未找到可导出的 Markdown 文件。")

    markdown_path = dataset_dir / validation.markdown_file
    markdown = markdown_path.read_text(encoding="utf-8")
    docx_path = _docx_output_path(export_base_dir, mode)

    # Pandoc converts inline LaTeX to native, editable Word math (OMML).
    # The legacy paragraph writer only handles standalone $$ blocks.
    if "$" in markdown and (pandoc_path := _find_pandoc()) is not None:
        result = subprocess.run(
            [
                str(pandoc_path), str(markdown_path),
                "--from=markdown+tex_math_dollars", "--to=docx",
                f"--resource-path={dataset_dir}", "--output", str(docx_path),
            ],
            capture_output=True, text=True, timeout=180,
        )
        if result.returncode != 0:
            raise ValueError(f"公式 Word 导出失败：{result.stderr.strip()}")
        document = Document(docx_path)
        _configure_document(document)
        for fonts in document.element.xpath(".//w:rFonts"):
            fonts.set(qn("w:eastAsia"), "PingFang SC")
        document.save(docx_path)
        return docx_path

    document = Document()
    _configure_document(document)
    _write_markdown_to_document(document, markdown, dataset_dir)

    document.save(docx_path)
    return docx_path


def _configure_document(document: Document) -> None:
    normal_style = document.styles["Normal"]
    normal_style.font.size = Pt(11)


def _export_pdf_layout_docx(
    export_base_dir: Path,
    source_pdf_path: Path,
    mode: str,
    page_texts: dict[int, str],
    fallback_markdown: str,
) -> Path:
    try:
        import fitz
    except ImportError as exc:
        raise ValueError("缺少 PyMuPDF 依赖，无法渲染 PDF 页面。") from exc
    return _export_facsimile_docx(
        export_base_dir, source_pdf_path, mode=mode,
        page_texts=page_texts, fallback_markdown=fallback_markdown,
    )


def _export_facsimile_docx(
    export_base_dir: Path,
    source_pdf_path: Path,
    *,
    mode: str = "facsimile",
    page_texts: dict[int, str] | None = None,
    fallback_markdown: str = "",
) -> Path:
    import fitz

    document = Document()
    _configure_document(document)
    with fitz.open(str(source_pdf_path)) as pdf_document, tempfile.TemporaryDirectory() as temp_dir:
        if pdf_document.page_count == 0:
            raise ValueError("PDF 没有页面。")
        section = document.sections[0]
        for page_index, page in enumerate(pdf_document):
            width, height = float(page.rect.width), float(page.rect.height)
            if page_index and (abs(width - section.page_width / EMUS_PER_POINT) > 0.5 or
                               abs(height - section.page_height / EMUS_PER_POINT) > 0.5):
                section = document.add_section(WD_SECTION_START.NEW_PAGE)
            elif page_index:
                document.add_page_break()
            section.page_width = Pt(width)
            section.page_height = Pt(height)
            section.left_margin = section.right_margin = Pt(0)
            section.top_margin = section.bottom_margin = Pt(0)
            section.header_distance = section.footer_distance = Pt(0)
            image_path = Path(temp_dir) / f"page_{page_index + 1:04d}.png"
            page.get_pixmap(matrix=fitz.Matrix(PDF_RENDER_DPI / 72, PDF_RENDER_DPI / 72), alpha=False).save(str(image_path))
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = Pt(1)
            paragraph.add_run().font.size = Pt(1)
            shape = paragraph.add_run().add_picture(str(image_path), width=Pt(width), height=Pt(height))
            _anchor_picture_to_page(shape, 0, 0, behind_text=True)

    if mode == "hybrid":
        document.add_page_break()
        document.add_heading("OCR 可编辑文本", level=1)
        if page_texts:
            for page_number, page_text in sorted(page_texts.items()):
                if page_text.strip():
                    _add_hybrid_page_text(document, page_number, page_text)
        elif fallback_markdown.strip():
            _add_ocr_text(document, _strip_markdown_images(fallback_markdown), font_size=9)

    docx_path = _docx_output_path(export_base_dir, mode)
    document.save(docx_path)
    return docx_path


def _export_layout_editable_docx(
    dataset_dir: Path, export_base_dir: Path, source_pdf_path: Path | None = None,
) -> Path:
    data = _load_dataset_json(dataset_dir)
    if data is None:
        raise ValueError("未找到可用于可编辑版式导出的 Docling JSON。")

    pages = data.get("pages") if isinstance(data, dict) else None
    if not isinstance(pages, dict) or not pages:
        raise ValueError("Docling JSON 中未找到页面尺寸信息。")

    document = Document()
    section = document.sections[0]
    first_page = pages.get(str(min(_sorted_page_numbers(pages)))) or {}
    first_size = first_page.get("size") or {}
    section.page_width = Pt(float(first_size.get("width") or 595.0))
    section.page_height = Pt(float(first_size.get("height") or 841.0))
    section.left_margin = section.right_margin = Pt(0)
    section.top_margin = section.bottom_margin = Pt(0)
    section.header_distance = section.footer_distance = Pt(0)
    _configure_document(document)

    elements_by_page = _layout_elements_by_page(data, dataset_dir)
    page_numbers = _sorted_page_numbers(pages)
    for page_index, page_number in enumerate(page_numbers):
        page_info = pages.get(str(page_number)) or {}
        page_size = page_info.get("size") if isinstance(page_info, dict) else {}
        page_width = float(page_size.get("width") or 595.0)
        page_height = float(page_size.get("height") or 841.0)

        if page_index and (abs(page_width - section.page_width / EMUS_PER_POINT) > 0.5 or
                           abs(page_height - section.page_height / EMUS_PER_POINT) > 0.5):
            section = document.add_section(WD_SECTION_START.NEW_PAGE)
            section.page_width = Pt(page_width)
            section.page_height = Pt(page_height)
            section.left_margin = section.right_margin = Pt(0)
            section.top_margin = section.bottom_margin = Pt(0)
            section.header_distance = section.footer_distance = Pt(0)
        elif page_index:
            document.add_page_break()

        # All images on the source page share a dedicated anchor paragraph.
        # Its position stays on this page even when editable text reflows.
        picture_elements = [element for element in elements_by_page.get(page_number, [])
                            if element["kind"] == "picture"]
        if picture_elements:
            anchor_paragraph = document.add_paragraph()
            anchor_paragraph.paragraph_format.space_before = Pt(0)
            anchor_paragraph.paragraph_format.space_after = Pt(0)
            anchor_paragraph.paragraph_format.line_spacing = Pt(1)
            anchor_paragraph.add_run().font.size = Pt(1)
            for element in picture_elements:
                _add_layout_picture(anchor_paragraph, element, page_width, page_height)

        for element in elements_by_page.get(page_number, []):
            if element["kind"] == "text":
                _add_layout_text(document, element, page_width, page_height,
                                 source_pdf_path=source_pdf_path, page_number=page_number)
        if not elements_by_page.get(page_number):
            # A terminal blank PDF page still needs a Word paragraph after
            # the page break, otherwise some renderers omit that page.
            paragraph = document.add_paragraph("\u00a0")
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.runs[0].font.size = Pt(1)

    docx_path = _docx_output_path(export_base_dir, "layout_editable")
    document.save(docx_path)
    return docx_path


def _load_dataset_json(dataset_dir: Path) -> dict | None:
    validation = validate_dataset(dataset_dir)
    json_candidates = []
    if validation.json_file:
        json_candidates.append(dataset_dir / validation.json_file)
    json_candidates.extend(sorted(dataset_dir.glob("*.json")))

    for json_path in _unique_paths(json_candidates):
        try:
            data = json.loads(json_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        if isinstance(data, dict) and isinstance(data.get("pages"), dict):
            return data
    return None


def _layout_elements_by_page(data: dict, dataset_dir: Path) -> dict[int, list[dict]]:
    pictures = data.get("pictures") if isinstance(data.get("pictures"), list) else []
    picture_boxes_by_page: dict[int, list[dict]] = {}
    elements_by_page: dict[int, list[dict]] = {}

    referenced_pictures = _referenced_picture_indices(data)
    for index, picture in enumerate(pictures):
        if not isinstance(picture, dict):
            continue
        if referenced_pictures is not None and index not in referenced_pictures:
            continue
        page_number, bbox = _entry_page_and_bbox(picture)
        if page_number is None or bbox is None:
            continue
        image_path = _resolve_picture_image_path(dataset_dir, picture, index)
        if image_path is None:
            continue
        picture_boxes_by_page.setdefault(page_number, []).append(bbox)
        elements_by_page.setdefault(page_number, []).append(
            {
                "kind": "picture",
                "bbox": bbox,
                "image_path": image_path,
                "index": index,
                "page": page_number,
            }
        )

    seen_texts: set[tuple] = set()
    texts = data.get("texts") if isinstance(data.get("texts"), list) else []
    for text_entry in texts:
        if not isinstance(text_entry, dict):
            continue
        page_number, bbox = _entry_page_and_bbox(text_entry)
        raw_text = str(text_entry.get("text") or text_entry.get("orig") or "").strip()
        if page_number is None or bbox is None or not raw_text:
            continue
        key = _layout_text_key(page_number, bbox, raw_text)
        if key in seen_texts:
            continue
        seen_texts.add(key)
        if _bbox_center_inside_any(bbox, picture_boxes_by_page.get(page_number, [])):
            continue

        elements_by_page.setdefault(page_number, []).append(
            {
                "kind": "text",
                "bbox": bbox,
                "text": raw_text,
                "label": text_entry.get("label") or "text",
                "marker": text_entry.get("marker") or "",
            }
        )

    for page_number, elements in elements_by_page.items():
        elements.sort(key=_layout_element_sort_key)
    return elements_by_page


def _referenced_picture_indices(data: dict) -> set[int] | None:
    """Skip picture objects removed from the document tree during OCR repairs."""
    body = data.get("body")
    if not isinstance(body, dict) or not isinstance(body.get("children"), list):
        return None
    found: set[int] = set()
    visited: set[str] = set()

    def visit(ref: str) -> None:
        if ref in visited or not ref.startswith("#/"):
            return
        visited.add(ref)
        parts = ref[2:].split("/")
        if len(parts) != 2:
            return
        kind, number = parts
        try:
            index = int(number)
            node = data[kind][index]
        except (KeyError, IndexError, TypeError, ValueError):
            return
        if kind == "pictures":
            found.add(index)
        for child in node.get("children", []):
            if isinstance(child, dict) and isinstance(child.get("$ref"), str):
                visit(child["$ref"])

    for child in body["children"]:
        if isinstance(child, dict) and isinstance(child.get("$ref"), str):
            visit(child["$ref"])
    return found


def _layout_element_sort_key(element: dict) -> tuple:
    bbox = element["bbox"]
    top = float(bbox["t"])
    left = float(bbox["l"])
    row_band = -math.floor(top / 18.0)
    return (row_band, left, 0 if element["kind"] == "text" else 1)


def _add_layout_text(
    document: Document,
    element: dict,
    page_width: float,
    page_height: float,
    *,
    source_pdf_path: Path | None = None,
    page_number: int | None = None,
) -> None:
    text = _clean_inline_markdown(element["text"])
    if not text:
        return
    marker = str(element.get("marker") or "").strip()
    if marker and not text.lstrip().startswith(marker):
        text = f"{marker} {text}"
    label = element.get("label") or "text"
    if label == "section_header":
        paragraph = document.add_heading(text, level=2)
    else:
        paragraph = document.add_paragraph()
    _anchor_text_frame(paragraph, element["bbox"], page_width, page_height)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0
    if label != "section_header":
        has_inline_math = bool(re.search(r"\$[^$]+\$", text))
        if label == "formula" and (text.startswith("\\") or not has_inline_math):
            math_element = _latex_to_omml(_normalize_latex_formula(text))
            if math_element is not None:
                paragraph._p.append(math_element)
            elif source_pdf_path is not None and page_number is not None and _add_formula_crop(
                paragraph, source_pdf_path, page_number, element["bbox"], page_width, page_height,
            ):
                pass
            else:
                paragraph.add_run(text).font.size = Pt(8.5)
        elif has_inline_math:
            _append_inline_math(paragraph, text)
        else:
            run = paragraph.add_run(text)
            run.font.size = Pt(8.5 if label in {"text", "list_item"} else 8)


def _anchor_text_frame(paragraph, bbox: dict, page_width: float, page_height: float) -> None:
    """Keep editable OCR text at its source-page coordinates."""
    rect = _picture_rect_points(bbox, page_width, page_height)
    if rect is None:
        return
    left, top, right, bottom = rect
    frame = OxmlElement("w:framePr")
    for name, value in {
        "hAnchor": "page", "vAnchor": "page", "wrap": "none", "anchorLock": "1",
        "x": str(round(left * 20)), "y": str(round(top * 20)),
        "w": str(round(min(page_width - left, max(right - left + 3, (right - left) * 1.06)) * 20)),
        "h": str(round(max(bottom - top + 3, 12) * 20)),
        "hRule": "atLeast", "hSpace": "0", "vSpace": "0",
    }.items():
        frame.set(qn(f"w:{name}"), value)
    paragraph._p.get_or_add_pPr().append(frame)


def _add_formula_crop(
    paragraph, source_pdf_path: Path, page_number: int, bbox: dict,
    page_width: float, page_height: float,
) -> bool:
    """Use the source pixels when damaged OCR cannot become native Word math."""
    rect = _picture_rect_points(bbox, page_width, page_height)
    if rect is None:
        return False
    left, top, right, bottom = rect
    try:
        import fitz

        with fitz.open(str(source_pdf_path)) as pdf_document:
            if not 1 <= page_number <= pdf_document.page_count:
                return False
            pixmap = pdf_document[page_number - 1].get_pixmap(
                matrix=fitz.Matrix(3, 3), clip=fitz.Rect(left, top, right, bottom), alpha=False,
            )
        with tempfile.TemporaryDirectory() as temporary:
            image_path = Path(temporary) / "formula.png"
            pixmap.save(str(image_path))
            shape = paragraph.add_run().add_picture(
                str(image_path), width=Pt(right - left), height=Pt(bottom - top),
            )
        shape._inline.docPr.set("descr", f"source-page={page_number};OCR-formula-fallback")
        return True
    except Exception:
        return False


def _append_inline_math(paragraph, text: str) -> None:
    cursor = 0
    for match in re.finditer(r"\$([^$]+)\$", text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor:match.start()]).font.size = Pt(8.5)
        math_element = _latex_to_omml(match.group(1).strip())
        inline = math_element.find(f".//{{{OMML_NAMESPACE}}}oMath") if math_element is not None else None
        if inline is not None:
            paragraph._p.append(inline)
        else:
            paragraph.add_run(match.group(0)).font.size = Pt(8.5)
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:]).font.size = Pt(8.5)


def _add_layout_picture(paragraph, element: dict, page_width: float, page_height: float) -> None:
    image_path = element["image_path"]
    geometry = _picture_rect_points(element["bbox"], page_width, page_height)
    if geometry is None:
        return
    left, top, right, bottom = geometry
    run = paragraph.add_run()
    try:
        shape = run.add_picture(str(image_path), width=Pt(right - left), height=Pt(bottom - top))
        shape._inline.docPr.set("descr", f"source-page={element['page']};picture-index={element['index']}")
        _anchor_picture_to_page(shape, left, top)
    except Exception:
        paragraph.add_run(f"[图片无法插入：{image_path.name}]")


def _picture_rect_points(bbox: dict, page_width: float, page_height: float) -> tuple[float, float, float, float] | None:
    left = max(0.0, float(bbox["l"]))
    right = min(page_width, float(bbox["r"]))
    if bbox.get("coord_origin") == "BOTTOMLEFT":
        top = page_height - float(bbox["t"])
        bottom = page_height - float(bbox["b"])
    else:
        top, bottom = float(bbox["t"]), float(bbox["b"])
    top = max(0.0, top)
    bottom = min(page_height, bottom)
    if right <= left or bottom <= top:
        return None
    return left, top, right, bottom


def _anchor_picture_to_page(shape, left_points: float, top_points: float, *, behind_text: bool = False) -> None:
    inline = shape._inline
    anchor = OxmlElement("wp:anchor")
    for name, value in {
        "distT": "0", "distB": "0", "distL": "0", "distR": "0",
        "simplePos": "0", "relativeHeight": "251658240", "behindDoc": "1" if behind_text else "0",
        "locked": "0", "layoutInCell": "1", "allowOverlap": "1",
    }.items():
        anchor.set(name, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    anchor.append(simple)
    for direction, offset_points in (("H", left_points), ("V", top_points)):
        position = OxmlElement(f"wp:position{direction}")
        position.set("relativeFrom", "page")
        offset = OxmlElement("wp:posOffset")
        offset.text = str(Pt(offset_points))
        position.append(offset)
        anchor.append(position)
    anchor.append(deepcopy(inline.extent))
    anchor.append(OxmlElement("wp:wrapNone"))
    anchor.append(deepcopy(inline.docPr))
    for frame_properties in inline.xpath("./wp:cNvGraphicFramePr"):
        anchor.append(deepcopy(frame_properties))
    anchor.append(deepcopy(inline.graphic))
    inline.getparent().replace(inline, anchor)


def _entry_page_and_bbox(entry: dict) -> tuple[int | None, dict | None]:
    prov = entry.get("prov")
    if not isinstance(prov, list) or not prov:
        return None, None
    first = prov[0]
    if not isinstance(first, dict):
        return None, None
    page_number = first.get("page_no")
    bbox = first.get("bbox")
    if not isinstance(page_number, int) or not isinstance(bbox, dict):
        return None, None
    required = {"l", "t", "r", "b"}
    if not required.issubset(bbox.keys()):
        return None, None
    return page_number, bbox


def _layout_text_key(page_number: int, bbox: dict, text: str) -> tuple:
    return (
        page_number,
        re.sub(r"\s+", "", text),
        round(float(bbox["l"]), 1),
        round(float(bbox["t"]), 1),
        round(float(bbox["r"]), 1),
        round(float(bbox["b"]), 1),
    )


def _bbox_center_inside_any(bbox: dict, candidates: list[dict]) -> bool:
    center_x = (float(bbox["l"]) + float(bbox["r"])) / 2
    center_y = (float(bbox["t"]) + float(bbox["b"])) / 2
    for candidate in candidates:
        if (
            float(candidate["l"]) <= center_x <= float(candidate["r"])
            and float(candidate["b"]) <= center_y <= float(candidate["t"])
        ):
            return True
    return False


def _resolve_picture_image_path(dataset_dir: Path, picture: dict, index: int) -> Path | None:
    uri = picture.get("image", {}).get("uri") if isinstance(picture.get("image"), dict) else ""
    if isinstance(uri, str) and uri:
        candidate = dataset_dir / "images" / Path(uri).name
        if candidate.exists() and candidate.is_file():
            return candidate

    matches = sorted((dataset_dir / "images").glob(f"image_{index:06d}_*"))
    if matches:
        return matches[0]
    return None


def _sorted_page_numbers(pages: dict) -> list[int]:
    page_numbers = []
    for key in pages:
        try:
            page_numbers.append(int(key))
        except (TypeError, ValueError):
            continue
    return sorted(page_numbers)


def _add_hybrid_page_text(document: Document, page_number: int, page_text: str) -> None:
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(2)
    run = heading.add_run(f"本页 OCR 文本（第 {page_number} 页）")
    run.bold = True
    run.font.size = Pt(9)
    _set_run_font(run, "Times New Roman")
    _add_ocr_text(document, page_text, font_size=8)


def _add_ocr_text(document: Document, text: str, font_size: int = 9) -> None:
    for block in _split_text_blocks(text):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        run = paragraph.add_run(block)
        run.font.size = Pt(font_size)
        _set_run_font(run, "Times New Roman")


def _write_markdown_to_document(document: Document, markdown: str, dataset_path: Path) -> None:
    paragraph_lines: list[str] = []
    formula_lines: list[str] = []
    in_formula = False

    def flush_paragraph() -> None:
        if not paragraph_lines:
            return
        text = " ".join(line.strip() for line in paragraph_lines if line.strip()).strip()
        paragraph_lines.clear()
        if text:
            document.add_paragraph(_clean_inline_markdown(text))

    for raw_line in markdown.splitlines():
        line = raw_line.rstrip()
        stripped = line.strip()

        if in_formula:
            if stripped.endswith(FORMULA_DELIMITER):
                formula_tail = stripped[: -len(FORMULA_DELIMITER)].strip()
                if formula_tail:
                    formula_lines.append(formula_tail)
                _add_formula(document, "\n".join(formula_lines).strip())
                formula_lines.clear()
                in_formula = False
            else:
                formula_lines.append(line)
            continue

        if not stripped:
            flush_paragraph()
            continue

        if stripped.startswith(FORMULA_DELIMITER):
            flush_paragraph()
            formula_body = stripped[len(FORMULA_DELIMITER) :].strip()
            if formula_body.endswith(FORMULA_DELIMITER):
                formula_body = formula_body[: -len(FORMULA_DELIMITER)].strip()
                _add_formula(document, formula_body)
            else:
                formula_lines = [formula_body] if formula_body else []
                in_formula = True
            continue

        heading_match = HEADING_PATTERN.match(line)
        if heading_match:
            flush_paragraph()
            level = min(len(heading_match.group(1)), 3)
            title = _clean_heading_text(heading_match.group(2))
            document.add_heading(title, level=level)
            continue

        if IMAGE_PATTERN.search(line):
            flush_paragraph()
            _write_line_with_images(document, line, dataset_path)
            continue

        bullet_match = BULLET_PATTERN.match(line)
        if bullet_match:
            flush_paragraph()
            _add_list_paragraph(document, bullet_match.group(1), "List Bullet")
            continue

        numbered_match = NUMBERED_PATTERN.match(line)
        if numbered_match:
            flush_paragraph()
            _add_list_paragraph(document, numbered_match.group(1), "List Number")
            continue

        paragraph_lines.append(line)

    if in_formula:
        _add_formula(document, "\n".join(formula_lines).strip())
    flush_paragraph()


def _write_line_with_images(document: Document, line: str, dataset_path: Path) -> None:
    cursor = 0
    for match in IMAGE_PATTERN.finditer(line):
        leading_text = line[cursor : match.start()].strip()
        if leading_text:
            document.add_paragraph(_clean_inline_markdown(leading_text))

        image_ref = match.group(2).strip()
        image_path = _resolve_image_path(dataset_path, image_ref)
        if image_path is None or not image_path.exists():
            document.add_paragraph(f"[图片缺失：{_display_image_ref(image_ref)}]")
        else:
            _add_image(document, image_path, _display_image_ref(image_ref))
        cursor = match.end()

    trailing_text = line[cursor:].strip()
    if trailing_text:
        document.add_paragraph(_clean_inline_markdown(trailing_text))


def _add_image(document: Document, image_path: Path, image_ref: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    try:
        run.add_picture(str(image_path), width=Inches(6.0))
    except Exception:
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
        paragraph.add_run(f"[图片无法插入：{image_ref}]")


def _add_formula(document: Document, formula: str) -> None:
    normalized_formula = _normalize_latex_formula(formula)
    if not normalized_formula:
        return

    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(3)

    omml = _latex_to_omml(normalized_formula)
    if omml is not None:
        paragraph._p.append(omml)
        return

    _add_latex_formula_text(paragraph, normalized_formula)


def _add_latex_formula_text(paragraph, formula: str) -> None:
    run = paragraph.add_run(f"$${formula}$$")
    run.font.size = Pt(12)
    _set_run_font(run, "Times New Roman")


def _normalize_latex_formula(formula: str) -> str:
    formula = html.unescape(str(formula or "")).strip()
    if formula.startswith(FORMULA_DELIMITER) and formula.endswith(FORMULA_DELIMITER):
        formula = formula[len(FORMULA_DELIMITER) : -len(FORMULA_DELIMITER)].strip()
    if formula.startswith(r"\[") and formula.endswith(r"\]"):
        formula = formula[2:-2].strip()
    formula = re.sub(r"\s+", " ", formula)
    formula = formula.replace("，", ",").replace("（", "(").replace("）", ")")
    formula = formula.replace("士", r"\pm")
    formula = re.sub(r"\s*\\,\s*", "", formula)
    formula = re.sub(r"\\(left|right)\s*([()\[\]{}|.])", r"\\\1\2", formula)
    formula = re.sub(r"\\([A-Za-z]+)\s+\{", r"\\\1{", formula)
    formula = re.sub(r"([_^])\s+\{", r"\1{", formula)
    formula = re.sub(r"\}\s+\{", r"}{", formula)
    formula = re.sub(r"\{\s+", "{", formula)
    formula = re.sub(r"\s+\}", "}", formula)
    formula = re.sub(r"\s+([_^=+\-*/(),])", r"\1", formula)
    formula = re.sub(r"([_^=+\-*/,(])\s+", r"\1", formula)
    formula = re.sub(r"\s+\)", ")", formula)
    formula = re.sub(r"\\quad\s*\(\s*(\d+)\s*-\s*(\d+)\s*\)", r"\\quad(\1-\2)", formula)
    formula = re.sub(r"\(\s*(\d+)\s*-\s*(\d+)\s*\)", r"(\1-\2)", formula)
    return formula.strip()


def _latex_to_omml(formula: str):
    pandoc_path = _find_pandoc()
    if pandoc_path is None:
        return None

    try:
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_path = Path(temp_dir)
            source_path = temp_path / "formula.md"
            docx_path = temp_path / "formula.docx"
            source_path.write_text(f"$${formula}$$\n", encoding="utf-8")
            subprocess.run(
                [
                    str(pandoc_path),
                    str(source_path),
                    "-f",
                    "markdown+tex_math_dollars",
                    "-t",
                    "docx",
                    "-o",
                    str(docx_path),
                ],
                check=True,
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                timeout=10,
            )
            with ZipFile(docx_path) as docx_file:
                document_xml = docx_file.read("word/document.xml")
            root = etree.fromstring(document_xml)
            omml_nodes = root.xpath(".//m:oMathPara", namespaces={"m": OMML_NAMESPACE})
            if not omml_nodes:
                omml_nodes = root.xpath(".//m:oMath", namespaces={"m": OMML_NAMESPACE})
            if not omml_nodes:
                return None
            return parse_xml(etree.tostring(omml_nodes[0], encoding="unicode"))
    except Exception:
        return None


def _find_pandoc() -> Path | None:
    discovered = shutil.which("pandoc")
    candidates = [
        Path(discovered) if discovered else None,
        Path("/usr/local/bin/pandoc"),
        Path("/opt/homebrew/bin/pandoc"),
    ]
    for candidate in candidates:
        if candidate is not None and candidate.exists() and candidate.is_file():
            return candidate
    return None


def _add_plain_formula(document: Document, formula: str) -> None:
    if not formula:
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(formula)
    run.font.size = Pt(12)
    _set_run_font(run, "Times New Roman")


def _looks_like_formula(text: str) -> bool:
    stripped = text.strip()
    return (
        stripped.startswith("$$")
        or stripped.endswith("$$")
        or bool(re.search(r"\\(?:frac|sqrt|sum|int|alpha|beta|theta|Delta)", stripped))
    )


def _add_list_paragraph(document: Document, text: str, style_name: str) -> None:
    try:
        document.add_paragraph(_clean_inline_markdown(text), style=style_name)
    except KeyError:
        document.add_paragraph(_clean_inline_markdown(text))


def _resolve_image_path(dataset_path: Path, image_ref: str) -> Path | None:
    image_ref = _display_image_ref(image_ref)
    parsed = urlparse(image_ref)
    if parsed.scheme or parsed.netloc:
        return None

    path_text = unquote(parsed.path)
    image_path = Path(path_text)
    if image_path.is_absolute():
        return image_path

    resolved = (dataset_path / image_path).resolve()
    try:
        resolved.relative_to(dataset_path)
    except ValueError:
        return None
    return resolved


def _display_image_ref(image_ref: str) -> str:
    image_ref = image_ref.strip()
    if image_ref.startswith("<") and image_ref.endswith(">"):
        image_ref = image_ref[1:-1].strip()
    match = re.match(r"^(\S+)(?:\s+['\"].*)?$", image_ref)
    if match:
        image_ref = match.group(1)
    for separator in ("#", "?"):
        if separator in image_ref:
            image_ref = image_ref.split(separator, 1)[0]
    return image_ref


def _clean_heading_text(text: str) -> str:
    text = re.sub(r"\s+#+\s*$", "", text).strip()
    return _clean_inline_markdown(text)


def _clean_inline_markdown(text: str) -> str:
    text = html.unescape(text)
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()


def _set_run_font(run, font_name: str) -> None:
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)


def _resolve_dataset_dir(raw_dataset_path: Path) -> Path | None:
    if raw_dataset_path.exists() and raw_dataset_path.is_dir():
        return raw_dataset_path
    if raw_dataset_path.exists() and raw_dataset_path.is_file():
        sibling_dataset_dir = raw_dataset_path.parent / raw_dataset_path.name
        if sibling_dataset_dir.exists() and sibling_dataset_dir.is_dir():
            return sibling_dataset_dir.resolve()
    return None


def _resolve_pdf_path(
    raw_dataset_path: Path,
    dataset_dir: Path | None,
    explicit_pdf_path: Path | None,
) -> Path | None:
    candidates: list[Path] = []
    if explicit_pdf_path is not None:
        explicit = explicit_pdf_path.expanduser().resolve()
        candidates.append(explicit)
        if explicit.exists() and explicit.is_dir():
            origin_name = _origin_pdf_name(explicit)
            if origin_name:
                candidates.append(explicit / origin_name)

    if raw_dataset_path.exists() and raw_dataset_path.is_file():
        candidates.append(raw_dataset_path)

    origin_name = _origin_pdf_name(dataset_dir) if dataset_dir is not None else None
    if dataset_dir is not None:
        if origin_name:
            candidates.extend(
                [
                    dataset_dir / origin_name,
                    dataset_dir.parent / origin_name,
                    dataset_dir.parent / f"{dataset_dir.name}.pdf",
                    dataset_dir.parent.parent / origin_name,
                ]
            )
        candidates.extend(
            [
                dataset_dir.parent / f"{_export_stem(dataset_dir)}.pdf",
                dataset_dir.parent / f"{dataset_dir.name}.pdf",
            ]
        )

    for candidate in _unique_paths(candidates):
        if candidate.exists() and candidate.is_file() and candidate.suffix.lower() == ".pdf":
            return candidate

    search_name = origin_name or (raw_dataset_path.name if raw_dataset_path.suffix.lower() == ".pdf" else "")
    if dataset_dir is not None and search_name:
        for root in _unique_paths([dataset_dir.parent, dataset_dir.parent.parent]):
            found = _find_pdf_by_name(root, search_name, dataset_dir)
            if found is not None:
                return found

    return None


def _origin_pdf_name(dataset_dir: Path | None) -> str:
    if dataset_dir is None or not dataset_dir.exists() or not dataset_dir.is_dir():
        return ""
    for json_path in sorted(dataset_dir.glob("*.json")):
        try:
            data = json.loads(json_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        origin = data.get("origin") if isinstance(data, dict) else None
        filename = origin.get("filename") if isinstance(origin, dict) else None
        if isinstance(filename, str) and filename.lower().endswith(".pdf"):
            return filename
    return ""


def _find_pdf_by_name(root: Path, filename: str, dataset_dir: Path) -> Path | None:
    if not root.exists() or not root.is_dir():
        return None
    matches = []
    for candidate in root.rglob(filename):
        if not candidate.is_file() or candidate.suffix.lower() != ".pdf":
            continue
        try:
            candidate.relative_to(dataset_dir)
            continue
        except ValueError:
            pass
        matches.append(candidate.resolve())
        if len(matches) >= 20:
            break
    return sorted(matches, key=lambda item: len(item.parts))[0] if matches else None


def _extract_page_texts(dataset_dir: Path | None) -> dict[int, str]:
    if dataset_dir is None or not dataset_dir.exists() or not dataset_dir.is_dir():
        return {}
    for json_path in sorted(dataset_dir.glob("*.json")):
        try:
            data = json.loads(json_path.read_text(encoding="utf-8"))
        except (OSError, json.JSONDecodeError):
            continue
        texts = data.get("texts") if isinstance(data, dict) else None
        if not isinstance(texts, list):
            continue
        page_texts: dict[int, list[str]] = {}
        for entry in texts:
            if not isinstance(entry, dict):
                continue
            text = str(entry.get("text") or entry.get("orig") or entry.get("content") or "").strip()
            page_number = _entry_page_number(entry)
            if text and page_number is not None:
                page_texts.setdefault(page_number, []).append(text)
        if page_texts:
            return {
                page_number: "\n".join(page_lines)
                for page_number, page_lines in sorted(page_texts.items())
            }
    return {}


def _entry_page_number(entry: dict) -> int | None:
    page_number = entry.get("page") or entry.get("page_no") or entry.get("page_number")
    if isinstance(page_number, int):
        return page_number
    prov = entry.get("prov")
    if isinstance(prov, list):
        for item in prov:
            if isinstance(item, dict) and isinstance(item.get("page_no"), int):
                return item["page_no"]
    return None


def _read_markdown_text(dataset_dir: Path | None) -> str:
    if dataset_dir is None:
        return ""
    validation = validate_dataset(dataset_dir)
    if not validation.valid or validation.markdown_file is None:
        return ""
    try:
        return (dataset_dir / validation.markdown_file).read_text(encoding="utf-8")
    except OSError:
        return ""


def _split_text_blocks(text: str) -> list[str]:
    return [block.strip() for block in re.split(r"\n{2,}", text) if block.strip()]


def _strip_markdown_images(markdown: str) -> str:
    lines = []
    for line in markdown.splitlines():
        if IMAGE_PATTERN.fullmatch(line.strip()):
            continue
        lines.append(_clean_inline_markdown(line))
    return "\n".join(lines)


def _unique_paths(paths: list[Path]) -> list[Path]:
    seen: set[str] = set()
    unique: list[Path] = []
    for path in paths:
        key = str(path)
        if key in seen:
            continue
        seen.add(key)
        unique.append(path)
    return unique


def _docx_output_path(export_base_dir: Path, mode: str) -> Path:
    exports_dir = export_base_dir / "exports"
    exports_dir.mkdir(parents=True, exist_ok=True)
    return exports_dir / f"{_export_stem(export_base_dir)}_{mode}.docx"


def _export_stem(dataset_path: Path) -> str:
    name_path = Path(dataset_path.name)
    return name_path.stem if name_path.suffix else name_path.name
